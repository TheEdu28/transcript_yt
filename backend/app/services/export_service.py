"""Exportación local y reproducible de recursos didácticos persistidos."""

import json
from pathlib import Path
from xml.sax.saxutils import escape

from app.core.config import Settings
from app.core.errors import PipelineError
from app.schemas.contracts import MaterialResponse


class ExportService:
    """Renderiza materiales estructurados a formatos portables."""

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def export(self, material: MaterialResponse, target_format: str, include_answers: bool = False) -> Path:
        """Write an export under the configured local data directory."""
        suffix_extra = "_con_respuestas" if include_answers else ""
        if target_format == "json":
            dumped = material.model_dump(mode="json")
            if not include_answers and material.material_type == "quiz":
                content = dumped.get("content", {})
                for q in content.get("multiple_choice", []):
                    q.pop("correct_option", None)
                    q.pop("explanation", None)
                for q in content.get("open_questions", []):
                    q.pop("expected_points", None)
            body, suffix = json.dumps(dumped, ensure_ascii=False, indent=2), "json"
        elif target_format == "markdown":
            body, suffix = self._as_markdown(material, include_answers=include_answers), "md"
        elif target_format == "docx":
            return self.export_to_docx(material, include_answers=include_answers)
        elif target_format == "pdf":
            return self.export_to_pdf(material, include_answers=include_answers)
        else:
            raise PipelineError("El formato de exportación no es compatible.", code="UNSUPPORTED_EXPORT_FORMAT")
        self.settings.export_directory.mkdir(parents=True, exist_ok=True)
        path = self.settings.export_directory / f"material_{material.id}{suffix_extra}.{suffix}"
        path.write_text(body, encoding="utf-8")
        return path

    def export_to_docx(self, material: MaterialResponse, include_answers: bool = False) -> Path:
        """Write a readable DOCX document from the validated material JSON."""
        try:
            from docx import Document
        except ImportError as error:
            raise PipelineError("La exportación DOCX requiere python-docx.", code="EXPORT_DEPENDENCY_MISSING", status_code=503) from error

        document = Document()
        content = material.content
        title_suffix = " (Con respuestas)" if include_answers and material.material_type == "quiz" else ""
        document.add_heading(self._title(material) + title_suffix, level=0)
        document.add_paragraph(f"Video: {material.video_id}")

        if material.material_type == "summary":
            document.add_heading("Sinopsis", level=1)
            document.add_paragraph(str(content["synopsis"]))
            document.add_heading("Glosario", level=1)
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, heading in zip(table.rows[0].cells, ("Término", "Definición", "Timestamp")):
                cell.text = heading
            for item in content["glossary"]:
                cells = table.add_row().cells
                cells[0].text = str(item["term"])
                cells[1].text = str(item["definition"])
                cells[2].text = str(item["timestamp"])
            document.add_heading("Bloques didácticos", level=1)
            for block in content["didactic_blocks"]:
                document.add_heading(str(block["title"]), level=2)
                document.add_paragraph(str(block["explanation"]))
                document.add_paragraph(f"Timestamps: {', '.join(block['timestamps'])}")
        else:
            self._add_quiz_docx(document, content, include_answers=include_answers)

        self.settings.export_directory.mkdir(parents=True, exist_ok=True)
        suffix_extra = "_con_respuestas" if include_answers else ""
        path = self.settings.export_directory / f"material_{material.id}{suffix_extra}.docx"
        document.save(path)
        return path

    def export_to_pdf(self, material: MaterialResponse, include_answers: bool = False) -> Path:
        """Write a readable PDF document from the validated material JSON."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as error:
            raise PipelineError("La exportación PDF requiere reportlab.", code="EXPORT_DEPENDENCY_MISSING", status_code=503) from error

        styles = getSampleStyleSheet()
        body = styles["BodyText"]
        heading = styles["Heading1"]
        subheading = styles["Heading2"]
        title_suffix = " (Con respuestas)" if include_answers and material.material_type == "quiz" else ""
        story = [Paragraph(escape(self._title(material) + title_suffix), styles["Title"]), Paragraph(f"Video: {material.video_id}", body), Spacer(1, 12)]
        content = material.content

        if material.material_type == "summary":
            story.extend([Paragraph("Sinopsis", heading), Paragraph(escape(str(content["synopsis"])), body), Spacer(1, 8)])
            story.append(Paragraph("Glosario", heading))
            rows = [[Paragraph("<b>Término</b>", body), Paragraph("<b>Definición</b>", body), Paragraph("<b>Timestamp</b>", body)]]
            rows.extend([
                [Paragraph(escape(str(item["term"])), body), Paragraph(escape(str(item["definition"])), body), Paragraph(escape(str(item["timestamp"])), body)]
                for item in content["glossary"]
            ])
            if len(rows) == 1:
                rows.append(["", "Sin términos.", ""])
            table = Table(rows, colWidths=[1.35 * inch, 4.2 * inch, 1.1 * inch], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eefc")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8c4d8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([table, Spacer(1, 12), Paragraph("Bloques didácticos", heading)])
            for block in content["didactic_blocks"]:
                story.extend([
                    Paragraph(escape(str(block["title"])), subheading),
                    Paragraph(escape(str(block["explanation"])), body),
                    Paragraph(f"Timestamps: {escape(', '.join(block['timestamps']))}", body),
                    Spacer(1, 8),
                ])
        else:
            self._add_quiz_pdf(story, content, heading, subheading, body, Paragraph, Spacer, include_answers=include_answers)

        self.settings.export_directory.mkdir(parents=True, exist_ok=True)
        suffix_extra = "_con_respuestas" if include_answers else ""
        path = self.settings.export_directory / f"material_{material.id}{suffix_extra}.pdf"
        SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.6 * inch, leftMargin=0.6 * inch).build(story)
        return path

    @staticmethod
    def _title(material: MaterialResponse) -> str:
        return f"Material didáctico {material.id} · {'Resumen' if material.material_type == 'summary' else 'Cuestionario'}"

    @staticmethod
    def _add_quiz_docx(document: object, content: dict, include_answers: bool = False) -> None:
        document.add_heading("Cuestionario", level=1)
        for index, question in enumerate(content.get("multiple_choice", []), start=1):
            document.add_heading(f"{index}. {question['question']}", level=2)
            for opt_idx, option in enumerate(question.get("options", [])):
                letter = chr(65 + opt_idx)
                document.add_paragraph(f"{letter}) {option}", style="List Bullet")
            document.add_paragraph(f"Evidencia: {question.get('evidence_timestamp', '')}")
            if include_answers:
                correct_idx = question.get("correct_option", 0)
                options = question.get("options", [])
                correct_val = options[correct_idx] if 0 <= correct_idx < len(options) else ""
                correct_letter = chr(65 + correct_idx) if 0 <= correct_idx < len(options) else ""
                document.add_paragraph(f"Respuesta correcta: {correct_letter}) {correct_val}")
                if question.get("explanation"):
                    document.add_paragraph(f"Explicación: {question['explanation']}")

        open_questions = content.get("open_questions", [])
        if open_questions:
            document.add_heading("Preguntas abiertas", level=1)
            for index, question in enumerate(open_questions, start=1):
                document.add_heading(f"{index}. {question['question']}", level=2)
                document.add_paragraph(f"Evidencia: {question.get('evidence_timestamp', '')}")
                if include_answers:
                    expected = question.get("expected_points", [])
                    document.add_paragraph(f"Puntos esperados: {', '.join(expected)}")
                else:
                    document.add_paragraph("Respuesta:")
                    document.add_paragraph("")

    @staticmethod
    def _add_quiz_pdf(story: list, content: dict, heading: object, subheading: object, body: object, paragraph: object, spacer: object, include_answers: bool = False) -> None:
        story.append(paragraph("Cuestionario", heading))
        for index, question in enumerate(content.get("multiple_choice", []), start=1):
            story.append(paragraph(f"{index}. {escape(str(question['question']))}", subheading))
            for opt_idx, option in enumerate(question.get("options", [])):
                letter = chr(65 + opt_idx)
                story.append(paragraph(f"• {letter}) {escape(str(option))}", body))
            story.append(paragraph(f"Evidencia: {escape(str(question.get('evidence_timestamp', '')))}", body))
            if include_answers:
                correct_idx = question.get("correct_option", 0)
                options = question.get("options", [])
                correct_val = options[correct_idx] if 0 <= correct_idx < len(options) else ""
                correct_letter = chr(65 + correct_idx) if 0 <= correct_idx < len(options) else ""
                story.append(paragraph(f"<b>Respuesta correcta:</b> {correct_letter}) {escape(str(correct_val))}", body))
                if question.get("explanation"):
                    story.append(paragraph(f"<b>Explicación:</b> {escape(str(question['explanation']))}", body))
            story.append(spacer(1, 8))

        open_questions = content.get("open_questions", [])
        if open_questions:
            story.append(paragraph("Preguntas abiertas", heading))
            for index, question in enumerate(open_questions, start=1):
                story.extend([
                    paragraph(f"{index}. {escape(str(question['question']))}", subheading),
                    paragraph(f"Evidencia: {escape(str(question.get('evidence_timestamp', '')))}", body),
                ])
                if include_answers:
                    expected = question.get("expected_points", [])
                    story.append(paragraph(f"<b>Puntos esperados:</b> {escape(', '.join(expected))}", body))
                else:
                    story.append(paragraph("<i>Espacio para responder:</i>", body))
                    story.append(spacer(1, 20))
                story.append(spacer(1, 8))

    @staticmethod
    def _as_markdown(material: MaterialResponse, include_answers: bool = False) -> str:
        """Render both supported resource types in a readable portable representation."""
        content = material.content
        title_suffix = " (Con respuestas)" if include_answers and material.material_type == "quiz" else ""
        lines = [f"# Material {material.id}{title_suffix}", "", f"Video: {material.video_id}", ""]
        if material.material_type == "summary":
            lines.extend(["## Sinopsis", "", str(content["synopsis"]), "", "## Glosario", ""])
            lines.extend(f"- **{item['term']}** ({item['timestamp']}): {item['definition']}" for item in content["glossary"])
            lines.extend(["", "## Bloques didácticos", ""])
            for block in content["didactic_blocks"]:
                lines.extend([f"### {block['title']}", "", str(block["explanation"]), "", f"Timestamps: {', '.join(block['timestamps'])}", ""])
        else:
            lines.extend(["## Cuestionario", ""])
            for index, question in enumerate(content.get("multiple_choice", []), start=1):
                lines.extend([f"### {index}. {question['question']}", ""])
                for opt_idx, option in enumerate(question.get("options", [])):
                    letter = chr(65 + opt_idx)
                    lines.append(f"- {letter}) {option}")
                lines.append(f"\nEvidencia: {question.get('evidence_timestamp', '')}")
                if include_answers:
                    correct_idx = question.get("correct_option", 0)
                    options = question.get("options", [])
                    correct_val = options[correct_idx] if 0 <= correct_idx < len(options) else ""
                    correct_letter = chr(65 + correct_idx) if 0 <= correct_idx < len(options) else ""
                    lines.append(f"**Respuesta correcta:** {correct_letter}) {correct_val}")
                    if question.get("explanation"):
                        lines.append(f"**Explicación:** {question['explanation']}")
                lines.append("")

            open_questions = content.get("open_questions", [])
            if open_questions:
                lines.extend(["## Preguntas abiertas", ""])
                for index, question in enumerate(open_questions, start=1):
                    lines.extend([f"### {index}. {question['question']}", "", f"Evidencia: {question.get('evidence_timestamp', '')}", ""])
                    if include_answers:
                        expected = question.get("expected_points", [])
                        if expected:
                            lines.extend([f"**Puntos esperados:** {', '.join(expected)}", ""])
        return "\n".join(lines).rstrip() + "\n"
